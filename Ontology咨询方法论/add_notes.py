import json, io, copy
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

BODY_XML = (
    '<p:sp xmlns:p="' + NS_P + '" xmlns:a="' + NS_A + '">'
    '<p:nvSpPr><p:cNvPr id="2" name="Notes Placeholder 1"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4343400"/>'
    '<a:ext cx="5486400" cy="4114800"/></a:xfrm></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/>'
    '<a:p><a:endParaRPr lang="zh-CN"/></a:p></p:txBody></p:sp>'
)


def ensure_body(ns):
    spTree = ns._element.find(qn('p:cSld')).find(qn('p:spTree'))
    for sp in spTree.findall(qn('p:sp')):
        nvPr = sp.find(qn('p:nvSpPr')).find(qn('p:nvPr'))
        ph = nvPr.find(qn('p:ph')) if nvPr is not None else None
        if ph is not None and ph.get('type') == 'body':
            return sp.find(qn('p:txBody'))
    sp = etree.fromstring(BODY_XML)
    spTree.append(sp)
    return sp.find(qn('p:txBody'))


def set_text(txBody, text):
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for para in text.split('\n'):
        p = etree.SubElement(txBody, qn('a:p'))
        if not para.strip():
            etree.SubElement(p, qn('a:endParaRPr')).set('lang', 'zh-CN')
            continue
        r = etree.SubElement(p, qn('a:r'))
        rPr = etree.SubElement(r, qn('a:rPr'))
        rPr.set('lang', 'zh-CN')
        rPr.set('sz', '1200')
        t = etree.SubElement(r, qn('a:t'))
        t.text = para


notes = json.load(io.open('notes.json', encoding='utf-8'))
prs = Presentation('Ontology本体咨询方法论.pptx')
print('slides', len(prs.slides), 'notes', len(notes))

for i, slide in enumerate(prs.slides):
    if i >= len(notes):
        break
    set_text(ensure_body(slide.notes_slide), notes[i])

prs.save('Ontology本体咨询方法论_最终.pptx')
print('saved')

# verify
import re, zipfile
z = zipfile.ZipFile('Ontology本体咨询方法论_最终.pptx')
ok = 0
chars = 0
mism = []
for i in range(1, 31):
    x = z.read('ppt/notesSlides/notesSlide%d.xml' % i).decode('utf-8')
    t = ''.join(re.findall(r'<a:t>(.*?)</a:t>', x, re.S))
    if t.strip():
        ok += 1
        chars += len(t)
        if t.strip()[:14] != notes[i-1].strip()[:14]:
            mism.append(i)
print('with_notes', ok, 'chars', chars, 'mismatch', mism)

prs2 = Presentation('Ontology本体咨询方法论_最终.pptx')
W, H = prs2.slide_width, prs2.slide_height
over = 0
for idx, s in enumerate(prs2.slides, 1):
    for sh in s.shapes:
        l, tp, w, h = sh.left, sh.top, sh.width, sh.height
        if l < -9144 or tp < -9144 or l + w > W + 9144 or tp + h > H + 9144:
            over += 1
            print('OVERFLOW p', idx, l, tp, w, h)
print('slides', len(prs2.slides), 'overflow', over)