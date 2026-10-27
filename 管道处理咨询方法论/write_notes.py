import json, re, shutil, zipfile
from pptx import Presentation

PPTX = '管道处理咨询方法论.pptx'
TMP = '_tmp_out.pptx'
notes = json.load(open('notes.json', encoding='utf-8'))

# 第一步：用 python-pptx 访问每页 notes_slide，触发 notesSlide part 生成
prs = Presentation(PPTX)
for slide in prs.slides:
    slide.notes_slide
prs.save(PPTX)
print('notes skeleton generated')

def build_sp(text, sid):
    parts = []
    for para in text.split('\n'):
        if not para.strip():
            parts.append('<a:p><a:endParaRPr lang="zh-CN"/></a:p>')
        else:
            esc = para.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            parts.append('<a:p><a:r><a:rPr lang="zh-CN" sz="1200" dirty="0"/><a:t>%s</a:t></a:r></a:p>' % esc)
    return (
        '<p:sp><p:nvSpPr>'
        '<p:cNvPr id="%d" name="Notes Placeholder %d"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr>'
        '</p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="685800" y="4343400"/><a:ext cx="5486400" cy="4114800"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        '<p:txBody><a:bodyPr wrap="square"><a:normAutofit/></a:bodyPr><a:lstStyle/>%s</p:txBody>'
        '</p:sp>' % (sid, sid, ''.join(parts))
    )

zin = zipfile.ZipFile(PPTX, 'r')
infos = zin.infolist()
names = zin.namelist()

slide_notes = {}
for rel in [n for n in names if re.match(r'ppt/slides/_rels/slide\d+\.xml\.rels$', n)]:
    sn = int(re.search(r'slide(\d+)\.xml\.rels$', rel).group(1))
    xml = zin.read(rel).decode('utf-8')
    m = re.search(r'Target="\.\./notesSlides/(notesSlide\d+\.xml)"', xml)
    if m:
        slide_notes[sn] = 'ppt/notesSlides/' + m.group(1)

zout = zipfile.ZipFile(TMP, 'w', zipfile.ZIP_DEFLATED)
written = 0
for item in infos:
    data = zin.read(item.filename)
    page = None
    for sn, nn in slide_notes.items():
        if nn == item.filename:
            page = sn
    if page:
        key = '%02d' % page
        if key in notes:
            xml = data.decode('utf-8')
            xml = re.sub(r'<p:sp><p:nvSpPr><p:cNvPr id="\d+" name="Notes Placeholder \d+"/>.*?</p:sp>', '', xml, flags=re.S)
            xml = xml.replace('</p:spTree>', build_sp(notes[key], 2) + '</p:spTree>')
            data = xml.encode('utf-8')
            written += 1
    zout.writestr(item, data)
zout.close(); zin.close()
shutil.move(TMP, PPTX)
print('notes parts written:', written)

z = zipfile.ZipFile(PPTX)
xmls = sorted([n for n in z.namelist() if re.match(r'ppt/notesSlides/notesSlide\d+\.xml$', n)],
              key=lambda s: int(re.search(r'notesSlide(\d+)', s).group(1)))
print('notes parts total:', len(xmls))
bad = 0; total = 0
for i, n in enumerate(xmls, 1):
    txt = ''.join(re.findall(r'<a:t>(.*?)</a:t>', z.read(n).decode('utf-8'), re.S))
    total += len(txt)
    for kw in ['讲述要点', '怎么讲', '时长', '本页', '这页', '要点一', '讲稿', '备注栏']:
        if kw in txt:
            print('VIOLATION p%d' % i, kw); bad += 1
    if '[' in txt or ']' in txt:
        print('BRACKET p%d' % i); bad += 1
    if len(txt) < 100:
        print('TOO SHORT p%d' % i, len(txt)); bad += 1
    if '\ufffd' in txt:
        print('ENCODING p%d' % i); bad += 1
print('total chars:', total, 'avg:', total // max(1, len(xmls)))
print('violations:', bad)

prs2 = Presentation(PPTX)
SW, SH = prs2.slide_width, prs2.slide_height
over = 0
for i, s in enumerate(prs2.slides, 1):
    for sh in s.shapes:
        try:
            if sh.left is None: continue
            if sh.left < -9525 or sh.top < -9525 or sh.left + sh.width > SW + 9525 or sh.top + sh.height > SH + 9525:
                over += 1
        except Exception:
            pass
print('slides:', len(prs2.slides), 'overflow:', over)
readable = sum(1 for s in prs2.slides
               if s.notes_slide.notes_text_frame is not None
               and len(s.notes_slide.notes_text_frame.text.strip()) > 50)
print('notes readable via python-pptx:', readable)
